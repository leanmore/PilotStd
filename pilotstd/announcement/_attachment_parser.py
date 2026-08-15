# 模块：项目//__解析器脚本
# 附件解析与正文提取—从解析器脚本拆分

import logging
import re
from io import BytesIO
from typing import Any, Optional

from bs4 import BeautifulSoup

from ._wps_utils import _clean_wps_fulltext, _split_wps_entries

logger = logging.getLogger(__name__)


def parse_attachment_text(attachment_bytes: bytes, filename: str = "") -> str:
    """根据附件文件名后缀选择合适的解析方法，提取纯文本。"""
    name_lower = filename.lower()
    if name_lower.endswith(".wps"):
        return parse_wps_text(attachment_bytes)
    if name_lower.endswith(".pdf"):
        return _parse_pdf_text(attachment_bytes)
    if name_lower.endswith((".doc", ".docx")):
        return _parse_docx_text(attachment_bytes)
    # 文件类型未知时依次尝试所有解析器，取首个有意义的输出
    for parser in (parse_wps_text, _parse_pdf_text, _parse_docx_text):
        try:
            text = parser(attachment_bytes)
            if text and len(text) > 20:
                return text
        except Exception:
            continue
    logger.warning("无法解析附件：未知格式")
    return ""


def parse_wps_text(raw_bytes: bytes) -> str:
    """从 .wps 文件的原始字节中提取可读文本。

    .wps 文件是 OLE2 容器内包 UTF-16LE 编码的文本。
    直接按 UTF-16LE 解码，清洗格式标记后按标准号切分条目。
    """
    from .parser import STD_CODE_PATTERN, _clean_wps_name

    try:
        text = raw_bytes.decode("utf-16-le", errors="ignore")
    except Exception:
        logger.debug("WPS 解码失败", exc_info=True)
        return ""

    text = _clean_wps_fulltext(text)
    entries = _split_wps_entries(text, STD_CODE_PATTERN)

    cleaned = []
    for entry in entries:
        # 去除不可打印字符，保留换行制表符
        chars = [ch for ch in entry if ch.isprintable() or ch in "\n\r\t"]
        entry = "".join(chars)
        entry = _clean_wps_name(entry)
        if entry:
            cleaned.append(entry)

    return "\n".join(cleaned)


def _parse_pdf_text(raw_bytes: bytes) -> str:
    """从 PDF 原始字节中提取纯文本（使用 pypdf）。"""
    from pypdf import PdfReader

    try:
        reader = PdfReader(BytesIO(raw_bytes))
        pages = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
        return "\n".join(pages)
    except Exception as e:
        logger.debug("PDF 解析失败: %s", e, exc_info=True)
        return ""


def _parse_docx_text(raw_bytes: bytes) -> str:
    """从 DOCX 原始字节中提取纯文本（使用 python-docx）。
    优先提取表格文本（公告标准清单通常在表格中），无表格则提取段落。
    """
    from docx import Document

    try:
        doc = Document(BytesIO(raw_bytes))
        lines = []
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(c for c in cells):
                    lines.append("\t".join(cells))
        if not lines:
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    lines.append(text)
        return "\n".join(lines)
    except Exception as e:
        # TODO(P2): .doc（OLE2）格式 python-docx 不支持，需另寻解析器（如 antiword/textract）或显式跳过标记
        logger.debug("DOCX 解析失败: %s", e, exc_info=True)
        return ""


def find_attachment_url(html: str) -> Optional[str]:
    """从公告详情页 HTML 中提取附件下载链接（支持 .wps / .docx / .doc / .pdf）。"""
    # 优先匹配标准公告附件链接，其次匹配通用文件下载链接
    for pattern in [
        r'href="(http://zxd\.sacinfo\.org\.cn/gb_notice/[^"]+)"',
        r'href="(https?://[^"]+\.(?:wps|docx?|pdf))"',
    ]:
        m = re.search(pattern, html, re.IGNORECASE)
        if m:
            return m.group(1)
    return None


def download_attachment(url: str, _http: Any = None) -> Optional[bytes]:
    """下载附件文件 (.wps)。返回原始字节，失败返回 None。"""
    from ..query.network import safe_raw_get

    if _http is not None:
        resp = _http.get(url)
        return resp.content if resp is not None and getattr(resp, "status_code", 0) == 200 else None
    resp = safe_raw_get(url, "announcement_attachment", timeout=60)
    if resp and resp.status_code == 200:
        return resp.content
    return None


def extract_content(html: str) -> str:
    """从公告详情页 HTML 中提取正文内容。DOM剪枝：先移除表格节点，再从剩余元素提取文本。"""
    if not html:
        return ""
    soup = BeautifulSoup(html, "html.parser")

    for table in soup.find_all("table"):
        table.decompose()

    # 提取所有段落文本；无段落时回退到关键词匹配的/元素
    all_p = soup.find_all("p")
    if all_p:
        return "\n\n".join(p.get_text(strip=True) for p in all_p)

    keywords = ["批准", "发布", "现予以", "公告如下"]
    candidates = [
        tag
        for tag in soup.find_all(["div", "p"])
        if any(kw in tag.get_text() for kw in keywords) and len(tag.get_text(strip=True)) > 50
    ]
    if candidates:
        best_match = max(candidates, key=lambda x: len(x.get_text()))
        return best_match.get_text(strip=True)

    return ""
