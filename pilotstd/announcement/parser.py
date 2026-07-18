# pilotstd/announcement/parser.py
# 公告数据解析 — 公告元数据提取 + HTML表格解析 + WPS附件文本提取

import logging
import re
from io import BytesIO
from typing import Any, Optional

from bs4 import BeautifulSoup

from ._wps_utils import _clean_wps_fulltext, _split_wps_entries

logger = logging.getLogger(__name__)

# 标准编号模式：代号 + 空格 + 顺序号[.部分号] + — + 年份
STD_CODE_PATTERN = re.compile(r"([A-Z]+(?:\d+)?(?:\s*/[A-Z]+)?)\s*(\d+(?:\.\d+)?)\s*[—\-\s]\s*(\d{4})")
# 代替标准在表格中的位置模式（"代替" 列的后面）
REPLACES_PATTERN = re.compile(r"([A-Z]+(?:\d+)?(?:\s*/[A-Z]+)?\s*\d+(?:\.\d+)?\s*[—\-]\s*\d{4})")


def parse_attachment_text(attachment_bytes: bytes, filename: str = "") -> str:
    """根据附件文件名后缀选择合适的解析方法，提取纯文本。"""
    name_lower = filename.lower()
    if name_lower.endswith(".wps"):
        return parse_wps_text(attachment_bytes)
    if name_lower.endswith(".pdf"):
        return _parse_pdf_text(attachment_bytes)
    if name_lower.endswith((".doc", ".docx")):
        return _parse_docx_text(attachment_bytes)
    # 后缀未知时尝试各解析器
    for parser in (parse_wps_text, _parse_pdf_text, _parse_docx_text):
        try:
            text = parser(attachment_bytes)
            if text and len(text) > 20:
                return text
        except Exception:
            continue
    logger.warning("无法解析附件：未知格式")
    return ""


def parse_wps_text(raw_bytes: bytes) -> str:
    """从 .wps 文件的原始字节中提取可读文本。

    .wps 文件是 OLE2 容器内包 UTF-16LE 编码的文本。
    直接按 UTF-16LE 解码，清洗格式标记后按标准号切分条目。
    """
    try:
        text = raw_bytes.decode("utf-16-le", errors="ignore")
    except Exception:
        logger.debug("WPS 解码失败", exc_info=True)
        return ""

    # 全文级清洗：去 WPS 格式标记、压缩空白
    text = _clean_wps_fulltext(text)

    # 按标准号位置切分粘连条目
    entries = _split_wps_entries(text, STD_CODE_PATTERN)

    # 逐条清洗字段级垃圾
    cleaned = []
    for entry in entries:
        # 过滤不可打印字符
        chars = [ch for ch in entry if ch.isprintable() or ch in "\n\r\t"]
        entry = "".join(chars)
        entry = _clean_wps_name(entry)
        if entry:
            cleaned.append(entry)

    return "\n".join(cleaned)


def _parse_pdf_text(raw_bytes: bytes) -> str:
    """从 PDF 原始字节中提取纯文本（使用 pypdf）。"""
    from pypdf import PdfReader

    try:
        reader = PdfReader(BytesIO(raw_bytes))
        pages = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
        return "\n".join(pages)
    except Exception as e:
        logger.debug("PDF 解析失败: %s", e, exc_info=True)
        return ""


def _parse_docx_text(raw_bytes: bytes) -> str:
    """从 DOCX 原始字节中提取纯文本（使用 python-docx）。
    优先提取表格文本（公告标准清单通常在表格中），无表格则提取段落。
    """
    from docx import Document

    try:
        doc = Document(BytesIO(raw_bytes))
        lines = []
        # 优先提取表格内容
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                # 过滤全空行
                if any(c for c in cells):
                    lines.append("\t".join(cells))
        if not lines:
            # 无表格时提取段落文本
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    lines.append(text)
        return "\n".join(lines)
    except Exception as e:
        logger.debug("DOCX 解析失败: %s", e, exc_info=True)
        return ""


def parse_announcement_meta(html: str) -> dict[str, str]:
    """从公告详情页 HTML 中提取公告级元数据。

    Returns:
        {title, publish_date}
    """
    meta = {"title": "", "publish_date": ""}
    soup = BeautifulSoup(html, "lxml")

    # 公告标题：<title> 标签内容
    title_tag = soup.find("title")
    if title_tag:
        meta["title"] = title_tag.get_text(strip=True)

    # 公告发布日期：排除表格内文本（实施日期在表格中），取正文区域第一个日期
    for table in soup.find_all("table"):
        table.decompose()
    text = soup.get_text()
    dates = re.findall(r"\d{4}-\d{2}-\d{2}", text)
    if dates:
        # 正文中第一个日期为落款日期（表格已排除，正文开头是发布/批准日期）
        meta["publish_date"] = dates[0]

    return meta


# 统一表头关键词 → 字段名映射（GB/HB/DB 三站点共用）
_HEADER_KEYWORD_MAP = [
    ("标准编号", "std_code"),
    ("编号", "std_code"),
    ("标准名称", "std_name"),
    ("名称", "std_name"),
    ("代替", "replaces_code"),
    ("实施日期", "implementation_date"),
    ("发布日期", "publish_date"),
    ("批准日期", "publish_date"),
    ("作废日期", "expiry_date"),
    ("废止日期", "expiry_date"),
    ("备案号", "record_no"),
    ("主管部门", "dept"),
]


def _build_header_map(table: Any) -> dict[int, str]:
    """从表格表头构建 {列索引: 字段名} 映射。"""
    col_map = {}
    thead = table.find("thead")
    headers = thead.find_all("th") if thead else []
    if not headers:
        first_row = table.find("tr")
        if first_row:
            headers = first_row.find_all("th")
    for i, h in enumerate(headers):
        # 去空格归一化（部分表格每个字间有空格，如"国 家 标 准"）
        h_text = h.get_text(strip=True).replace(" ", "")
        for keyword, field_name in _HEADER_KEYWORD_MAP:
            if keyword in h_text:
                col_map[i] = field_name
                break
    return col_map


def _find_col(col_map: dict[int, str], field_name: str) -> Optional[int]:
    """在 col_map 中查找字段名对应的列索引。找不到返回 None。"""
    for idx, name in col_map.items():
        if name == field_name:
            return idx
    return None


def parse_html_table(html: str) -> list[dict[str, Any]]:
    """从公告详情页 HTML 表格中提取标准列表。表头驱动列识别，兼容 GB(5列)/HB(8列)/DB(8列)。
    遍历所有表格，返回第一个同时含标准编号和标准名称列的数据表格。

    Returns:
        [{std_code, std_name, replaces_code, publish_date, implementation_date, record_no, dept}, ...]
    """
    soup = BeautifulSoup(html, "lxml")

    # 遍历所有表格，找第一个含标准编号+标准名称列的数据表格
    for table in soup.find_all("table"):
        col_map = _build_header_map(table)
        code_col = _find_col(col_map, "std_code")
        name_col = _find_col(col_map, "std_name")
        if code_col is None or name_col is None:
            continue

        results = []
        tbody = table.find("tbody") or table
        for tr in tbody.find_all("tr"):
            cells = tr.find_all("td")
            if len(cells) < max(code_col, name_col) + 1:
                continue
            cell_texts = [c.get_text(strip=True) for c in cells]
            # 首列非数字序号 → 跳过（非数据行）
            if not cell_texts[0].isdigit():
                continue

            std_code = cell_texts[code_col]
            if not STD_CODE_PATTERN.match(std_code):
                continue

            item = {
                "std_code": std_code,
                "std_name": cell_texts[name_col] if name_col < len(cell_texts) else "",
                "replaces_code": "",
                "publish_date": "",
                "implementation_date": "",
                "expiry_date": "",
                "record_no": "",
                "dept": "",
            }
            for col_idx, field_name in col_map.items():
                if col_idx < len(cell_texts) and field_name in item:
                    item[field_name] = cell_texts[col_idx]
            results.append(item)

        return results

    return []


def _clean_wps_name(name: str) -> str:
    """清理 WPS 提取文本中的二进制垃圾和非表头残留。"""
    # WPS 页脚标记（安全网：_clean_wps_fulltext 遗漏的在这里补刀）
    name = re.sub(r"\s*—+\s*PAGE\s*\n?\s*MERGEFORMAT\s*\d*\s*—*\s*", " ", name, flags=re.IGNORECASE)
    name = re.sub(r"PAGE\s*\d+\s*OF\s*\d+", " ", name, flags=re.IGNORECASE)
    # 去除不可打印字符和控制字符（保留中英文、数字、常见标点）
    name = re.sub(
        r"[^一-鿿　-〿＀-￯"
        r"a-zA-Z0-9\s\-—/\.\(\)（）\d]",
        "",
        name,
    )
    # 去除 WPS 表格表头关键词残留
    for kw in [
        "国家标准",
        "行业标准",
        "地方标准",
        "指导性技术文件",
        "标准化指导性技术文件编号",
        "标准化指导性技术文件名称",
        "标准编号",
        "标准名称",
        "代替标准号",
        "实施日期",
        "发布日期",
        "复审结论",
        "标准废止日期",
        "代替文件号",
    ]:
        name = name.replace(kw, "")
    return name.strip()


def _find_field_text(text: str, matches: list[re.Match[str]], i: int, skip_indices: set[int]) -> tuple[str, str, int]:
    """当前匹配到下一个匹配之间的字段文本 + 代替号检测。返回 (field_text, replaces_code, next_idx)。"""
    next_idx = i + 1
    while next_idx < len(matches) and next_idx in skip_indices:
        next_idx += 1

    if next_idx < len(matches):
        field_text = text[matches[i].end() : matches[next_idx].start()]
    else:
        field_text = text[matches[i].end() : matches[i].end() + 300]

    replaces_code = ""
    if next_idx < len(matches):
        has_chinese = bool(re.search(r"[一-鿿]", field_text))
        if has_chinese and "\n" not in field_text:
            replaces_code = matches[next_idx].group(0)
            skip_indices.add(next_idx)
            post_replaces = text[matches[next_idx].end() : matches[next_idx].end() + 50]
            field_text = field_text + post_replaces

    return field_text, replaces_code, next_idx


def _parse_entry_fields(field_text: str) -> tuple[str, str, str]:
    """从字段文本中提取标准名称、代替号、发布日期。返回 (std_name, replaces_code, publish_date)。"""
    replaces_code = ""
    std_name = _clean_wps_name(field_text)

    replaces_match = REPLACES_PATTERN.search(field_text)
    if replaces_match:
        rep_start = replaces_match.start()
        std_name = _clean_wps_name(field_text[:rep_start])
        replaces_code = replaces_match.group(0)
        field_text = field_text[replaces_match.end() :]

    date_match = re.search(r"(\d{4}-\d{2}-\d{2})", field_text)
    publish_date = date_match.group(1) if date_match else ""

    return std_name, replaces_code, publish_date


def parse_text_table(text: str) -> list[dict[str, Any]]:
    """从附件文本中提取标准表格。不依赖换行符，用标准编号模式全文本扫描。

    兼容多种 WPS/PDF 提取格式：
      - 标准逐行（每行一条）  → 旧格式，仍支持
      - 同一条目跨行（表头折行） → STD_CODE_PATTERN 跨行不依赖 \\n
      - 多条目拼接在一行（WPS 二进制提取） → finditer 扫描全文本
      - 二进制垃圾前缀（WPS 格式头） → 正则扫描跳过非标准号内容

    Returns:
        [{std_code, std_name, replaces_code, publish_date}, ...]
    """
    results: list[dict[str, Any]] = []
    matches = list(STD_CODE_PATTERN.finditer(text))
    if not matches:
        return results

    skip_indices: set[int] = set()

    for i, match in enumerate(matches):
        if i in skip_indices:
            continue
        std_code = match.group(0)

        field_text, replaces_code_adjacent, _ = _find_field_text(text, matches, i, skip_indices)

        if not replaces_code_adjacent:
            replaces_code_adjacent = ""  # _find_field_text may return ""

        std_name, replaces_code, publish_date = _parse_entry_fields(field_text)
        # 如果 _find_field_text 检测到了代替号但 REPLACES_PATTERN 没扫到，用前者
        if not replaces_code and replaces_code_adjacent:
            replaces_code = replaces_code_adjacent

        if len(std_name) < 2:
            continue

        results.append(
            {
                "std_code": std_code,
                "std_name": std_name,
                "replaces_code": replaces_code,
                "publish_date": publish_date,
                "implementation_date": "",
            }
        )

    return results


def _code_key(item: dict[str, Any]) -> str:
    """标准号去重键：代号 + 名称前20字符，用于 HTML 与附件交叉去重。"""
    return item.get("std_code", "") + "|" + item.get("std_name", "")[:20]  # type: ignore[no-any-return]  # dict.get 返回 Any


def _ocr_pdf(pdf_bytes: bytes, ocr_provider: Any) -> str:
    """用 OCR 提供商识别 PDF 全部页面，返回合并文本。"""
    # OcrScheduler 内部已拆页+调度，直接返回全文，无需逐页循环
    from pilotstd.announcement.ocr import OcrScheduler

    if isinstance(ocr_provider, OcrScheduler):
        result = ocr_provider.recognize_pdf(pdf_bytes, page_num=1)
        return result or ""

    total_pages = 1
    try:
        from io import BytesIO

        from pypdf import PdfReader

        reader = PdfReader(BytesIO(pdf_bytes))
        total_pages = len(reader.pages)
    except Exception:
        logger.debug("PDF 页数检测失败，假定 1 页", exc_info=True)

    texts = []
    failed = 0
    for page_num in range(1, total_pages + 1):
        result = ocr_provider.recognize_pdf(pdf_bytes, page_num=page_num)
        if result is not None and getattr(result, "ok", False):  # type: ignore[union-attr]
            texts.append(result.text)  # type: ignore[union-attr]
        else:
            failed += 1

    result = "\n".join(texts)
    logger.info(
        "OCR 完成: %d/%d 页成功, %d 字符",
        total_pages - failed,
        total_pages,
        len(result),
    )
    return result


def parse_announcement_detail(
    html: str,
    attachment_bytes: Optional[bytes] = None,
    attachment_filename: str = "",
    ocr_provider: Any = None,
) -> tuple[list[dict[str, Any]], dict[str, str]]:
    """解析公告详情页，HTML + 附件交叉校验补全。

    三层回退：
      1. HTML 表格解析（覆盖 ~99.5% 公告）
      2. 附件解析：.wps/.docx 本地、.pdf 先 pypdf 再 OCR
      3. HTML + 附件交叉去重合并

    Args:
        html: 详情页 HTML
        attachment_bytes: 附件原始字节（可选）
        attachment_filename: 附件文件名或URL，用于判断格式（.wps/.pdf/.docx）
        ocr_provider: 可选，BaseOcrProvider 实例，用于 PDF OCR 识别

    Returns:
        (items, meta)
    """
    meta = parse_announcement_meta(html)
    html_items = parse_html_table(html) or []

    # HTML 缺发布日期时用公告落款日期补
    for item in html_items:
        if not item.get("publish_date") and meta["publish_date"]:
            item["publish_date"] = meta["publish_date"]

    # 附件解析
    att_items = []
    if attachment_bytes:
        text = parse_attachment_text(attachment_bytes, filename=attachment_filename)
        if text:
            att_items = parse_text_table(text)
        # PDF 文本提取失败 → 尝试 OCR
        elif attachment_filename.lower().endswith(".pdf") or (
            attachment_filename and ".pdf" in attachment_filename.lower()
        ):
            if ocr_provider:
                att_text = _ocr_pdf(attachment_bytes, ocr_provider)
                if att_text:
                    att_items = parse_text_table(att_text)

    # 交叉去重：HTML 为主，附件补充未覆盖的标准号
    if html_items and att_items:
        codes = {_code_key(i) for i in html_items}
        new_count = 0
        for item in att_items:
            if _code_key(item) not in codes:
                html_items.append(item)
                new_count += 1
        logger.info("HTML %d 条 + 附件补充 %d 条", len(html_items) - new_count, new_count)
        return html_items, meta

    if html_items:
        logger.info("HTML 表格解析到 %d 条标准", len(html_items))
        return html_items, meta

    if att_items:
        logger.info("附件解析到 %d 条标准", len(att_items))
        return att_items, meta

    title = meta.get("title", "") or ""
    logger.warning("公告详情页未能解析出标准列表: %s", title[:80])
    return [], meta


def find_attachment_url(html: str) -> Optional[str]:
    """从公告详情页 HTML 中提取附件下载链接（支持 .wps / .docx / .doc / .pdf）。"""
    # zxd.sacinfo.org.cn 附件服务器
    for pattern in [
        r'href="(http://zxd\.sacinfo\.org\.cn/gb_notice/[^"]+)"',
        r'href="(https?://[^"]+\.(?:wps|docx?|pdf))"',
    ]:
        m = re.search(pattern, html, re.IGNORECASE)
        if m:
            return m.group(1)
    return None


def download_attachment(url: str, _http: Any = None) -> Optional[bytes]:
    """下载附件文件 (.wps)。返回原始字节，失败返回 None。
    _http: DI 注入，可传入 mock HTTP 客户端，默认 None 时使用 safe_raw_get。
    """
    from ..query.network import safe_raw_get

    if _http is not None:
        resp = _http.get(url)
        return resp.content if resp is not None and getattr(resp, "status_code", 0) == 200 else None
    resp = safe_raw_get(url, "announcement_attachment", timeout=60)
    if resp and resp.status_code == 200:
        return resp.content
    return None


def extract_content(html: str) -> str:
    """从公告详情页 HTML 中提取正文内容。三层回退：精确选择器 → p标签 → 关键词启发式。"""
    if not html:
        return ""
    soup = BeautifulSoup(html, "html.parser")

    # 1. 优先使用精确选择器
    content_selectors = [
        "div.content",
        "div.article-content",
        "div.main-text",
        "div#Content",
        ".announcement-body",
    ]
    for selector in content_selectors:
        element = soup.select_one(selector)
        if element:
            paragraphs = element.find_all("p")
            if paragraphs:
                return "\n\n".join(p.get_text(strip=True) for p in paragraphs)
            return element.get_text(strip=True)

    # 2. 所有 p 标签（排除表格内）
    all_p = soup.find_all("p")
    if all_p:
        return "\n\n".join(p.get_text(strip=True) for p in all_p if not p.find_parent("table"))

    # 3. 启发式兜底：查找含关键词且长度 > 50 的块级元素
    keywords = ["批准", "发布", "现予以", "公告如下"]
    candidates = [
        tag
        for tag in soup.find_all(["div", "p"])
        if any(kw in tag.get_text() for kw in keywords) and len(tag.get_text(strip=True)) > 50
    ]
    if candidates:
        best_match = max(candidates, key=lambda x: len(x.get_text()))
        return best_match.get_text(strip=True)

    return ""
